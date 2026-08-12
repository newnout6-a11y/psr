"""Fail-closed attachment parsing and safe model-context construction."""

from __future__ import annotations

from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from enum import StrEnum
from hashlib import sha256
from io import BytesIO, StringIO
import csv
import json
import math
import re
from typing import Any
import unicodedata
import zipfile


class BuyerAttachmentType(StrEnum):
    TEXT = "text"
    JSON = "json"
    CSV = "csv"
    PDF = "pdf"
    DOCX = "docx"
    XLSX = "xlsx"
    PPTX = "pptx"
    XML = "xml"
    IMAGE = "image"
    ZIP = "zip"


class BuyerAttachmentParseStatus(StrEnum):
    PARSED = "parsed"
    UNSUPPORTED = "unsupported"
    TOO_LARGE = "too_large"
    CORRUPT = "corrupt"


_EXTENSION_TYPES: dict[str, BuyerAttachmentType] = {
    ".txt": BuyerAttachmentType.TEXT,
    ".md": BuyerAttachmentType.TEXT,
    ".log": BuyerAttachmentType.TEXT,
    ".json": BuyerAttachmentType.JSON,
    ".csv": BuyerAttachmentType.CSV,
    ".pdf": BuyerAttachmentType.PDF,
    ".docx": BuyerAttachmentType.DOCX,
    ".xlsx": BuyerAttachmentType.XLSX,
    ".pptx": BuyerAttachmentType.PPTX,
    ".xml": BuyerAttachmentType.XML,
    ".svg": BuyerAttachmentType.XML,
    ".png": BuyerAttachmentType.IMAGE,
    ".jpg": BuyerAttachmentType.IMAGE,
    ".jpeg": BuyerAttachmentType.IMAGE,
    ".webp": BuyerAttachmentType.IMAGE,
    ".gif": BuyerAttachmentType.IMAGE,
    ".zip": BuyerAttachmentType.ZIP,
}
_MIME_TYPES: dict[str, BuyerAttachmentType] = {
    "text/plain": BuyerAttachmentType.TEXT,
    "text/markdown": BuyerAttachmentType.TEXT,
    "application/json": BuyerAttachmentType.JSON,
    "text/json": BuyerAttachmentType.JSON,
    "text/csv": BuyerAttachmentType.CSV,
    "application/csv": BuyerAttachmentType.CSV,
    "application/pdf": BuyerAttachmentType.PDF,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": BuyerAttachmentType.DOCX,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": BuyerAttachmentType.XLSX,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": BuyerAttachmentType.PPTX,
    "application/xml": BuyerAttachmentType.XML,
    "text/xml": BuyerAttachmentType.XML,
    "image/png": BuyerAttachmentType.IMAGE,
    "image/jpeg": BuyerAttachmentType.IMAGE,
    "image/webp": BuyerAttachmentType.IMAGE,
    "image/gif": BuyerAttachmentType.IMAGE,
    "application/zip": BuyerAttachmentType.ZIP,
    "application/x-zip-compressed": BuyerAttachmentType.ZIP,
}
_GENERIC_MIME_TYPES = {"application/octet-stream", "binary/octet-stream"}
_INVALID_FILENAME_CHARS = frozenset('<>:"/\\|?*')
_WINDOWS_RESERVED_FILENAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{index}" for index in range(1, 10)),
    *(f"LPT{index}" for index in range(1, 10)),
}
_MIME_RE = re.compile(r"^[a-z0-9!#$&^_.+-]+/[a-z0-9!#$&^_.+-]+$")
_SENSITIVE_KEY_RE = re.compile(
    r"(?i)(?:^|[_-])(?:authorization|cookie|password|secret|token|api[_-]?key|session(?:id)?)(?:$|[_-])"
)
_SECRET_HEADER_RE = re.compile(
    r"(?im)^(?P<prefix>\s*(?:authorization|cookie|set-cookie|x-api-key|api[_-]?key|"
    r"access[_-]?token|refresh[_-]?token|token|secret|password|session(?:[_-]?id)?)\s*[:=]).*$"
)
_SECRET_ASSIGNMENT_RE = re.compile(
    r"(?ix)(?P<prefix>[\"']?(?:authorization|cookie|set-cookie|x-api-key|api[_-]?key|"
    r"access[_-]?token|refresh[_-]?token|token|secret|password|session(?:[_-]?id)?)[\"']?\s*[:=]\s*)"
    r"(?P<value>\"(?:[^\"\\]|\\.)*\"|'(?:[^'\\]|\\.)*'|[^,;\s}\]]+)"
)
_BEARER_RE = re.compile(r"(?i)\bbearer\s+[a-z0-9._~+/=-]+")
_URL_SECRET_RE = re.compile(
    r"(?i)(?P<prefix>[?&](?:authorization|cookie|api[_-]?key|access[_-]?token|refresh[_-]?token|token|secret|password|session(?:[_-]?id)?)=)[^&#\s]+"
)


@dataclass(frozen=True, slots=True)
class BuyerAttachmentParseLimits:
    """Bound every parser before attachment content reaches the model context."""

    max_bytes: int = 10 * 1024 * 1024
    max_characters: int = 24_000
    max_rows: int = 200
    max_columns: int = 50
    max_pages: int = 20
    max_sheets: int = 10
    max_slides: int = 20
    max_archive_members: int = 100

    def __post_init__(self) -> None:
        for name in (
            "max_bytes",
            "max_characters",
            "max_rows",
            "max_columns",
            "max_pages",
            "max_sheets",
            "max_slides",
            "max_archive_members",
        ):
            value = getattr(self, name)
            if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
                raise ValueError(f"{name} must be a positive integer")


@dataclass(frozen=True, slots=True)
class BuyerAttachmentMetadata:
    """Immutable metadata recorded before any parser is selected."""

    filename: str
    content_type: str | None
    attachment_type: BuyerAttachmentType | None
    size_bytes: int
    sha256: str

    @property
    def supported(self) -> bool:
        return self.attachment_type is not None

    def to_manifest(self) -> dict[str, Any]:
        return {
            "filename": self.filename,
            "content_type": self.content_type,
            "attachment_type": self.attachment_type.value if self.attachment_type is not None else None,
            "size_bytes": self.size_bytes,
            "sha256": self.sha256,
        }


@dataclass(frozen=True, slots=True)
class BuyerAttachmentOmission:
    """An explicit reason that content was not included or fully processed."""

    reason: str
    detail: str | None = None

    def to_manifest(self) -> dict[str, str | None]:
        return {"reason": self.reason, "detail": self.detail}


@dataclass(frozen=True, slots=True)
class BuyerAttachmentParseResult:
    """Parser output that can safely be persisted and shown to an operator."""

    metadata: BuyerAttachmentMetadata
    status: BuyerAttachmentParseStatus
    parser: str | None = None
    text: str = ""
    omissions: tuple[BuyerAttachmentOmission, ...] = ()

    def to_manifest(self) -> dict[str, Any]:
        manifest = self.metadata.to_manifest()
        manifest.update(
            {
                "status": self.status.value,
                "parser": self.parser,
                "extracted_characters": len(self.text),
                "omissions": [omission.to_manifest() for omission in self.omissions],
            }
        )
        return manifest


@dataclass(frozen=True, slots=True)
class BuyerAttachmentContext:
    """Bounded text and manifest suitable for a proposal or scoring prompt."""

    text: str
    manifest: tuple[dict[str, Any], ...]
    context_hash: str


def normalize_attachment_filename(value: str, *, fallback: str = "attachment", max_length: int = 160) -> str:
    """Return a display-safe basename that cannot express a local path."""

    if not isinstance(value, str):
        raise TypeError("filename must be a string")
    if not isinstance(fallback, str) or not fallback.strip():
        raise ValueError("fallback must be a non-empty string")
    if isinstance(max_length, bool) or not isinstance(max_length, int) or max_length < 16:
        raise ValueError("max_length must be an integer of at least 16")

    name = unicodedata.normalize("NFKC", value).replace("\\", "/").rsplit("/", 1)[-1]
    name = "".join(character for character in name if ord(character) >= 32 and character not in _INVALID_FILENAME_CHARS)
    name = " ".join(name.split()).strip(" .")
    if not name or name in {".", ".."}:
        name = fallback.strip()

    stem, extension = _filename_parts(name)
    if stem.upper() in _WINDOWS_RESERVED_FILENAMES:
        stem = f"{stem}_"
    extension = extension[:32]
    available_stem = max_length - len(extension)
    if available_stem <= 0:
        extension = ""
        available_stem = max_length
    stem = stem[:available_stem].rstrip(" .") or fallback.strip()[:available_stem]
    return f"{stem}{extension}"


def normalize_attachment_content_type(value: str | None) -> str | None:
    """Return a valid MIME value without parameters, or ``None`` if malformed."""

    if value is None:
        return None
    if not isinstance(value, str):
        raise TypeError("content_type must be a string or None")
    content_type = value.split(";", 1)[0].strip().casefold()
    return content_type if _MIME_RE.fullmatch(content_type) else None


def detect_attachment_type(filename: str, content_type: str | None = None) -> BuyerAttachmentType | None:
    """Detect only a known, allowlisted attachment type; reject conflicting hints."""

    normalized_filename = normalize_attachment_filename(filename)
    _stem, extension = _filename_parts(normalized_filename)
    extension_type = _EXTENSION_TYPES.get(extension.casefold())
    normalized_content_type = normalize_attachment_content_type(content_type)
    mime_type = _MIME_TYPES.get(normalized_content_type or "")
    if mime_type is not None and extension_type is not None and mime_type is not extension_type:
        return None
    if mime_type is not None:
        return mime_type
    if normalized_content_type is None or normalized_content_type in _GENERIC_MIME_TYPES:
        return extension_type
    return None


def build_attachment_metadata(
    content: bytes | bytearray | memoryview,
    *,
    filename: str,
    content_type: str | None = None,
) -> BuyerAttachmentMetadata:
    """Compute the immutable metadata used for dedupe and parser dispatch."""

    payload = _as_bytes(content)
    normalized_filename = normalize_attachment_filename(filename)
    normalized_content_type = normalize_attachment_content_type(content_type)
    return BuyerAttachmentMetadata(
        filename=normalized_filename,
        content_type=normalized_content_type,
        attachment_type=detect_attachment_type(normalized_filename, normalized_content_type),
        size_bytes=len(payload),
        sha256=sha256(payload).hexdigest(),
    )


def parse_attachment(
    content: bytes | bytearray | memoryview,
    *,
    filename: str,
    content_type: str | None = None,
    limits: BuyerAttachmentParseLimits | None = None,
) -> BuyerAttachmentParseResult:
    """Parse one allowlisted attachment without executing its contents."""

    active_limits = limits or BuyerAttachmentParseLimits()
    if not isinstance(active_limits, BuyerAttachmentParseLimits):
        raise TypeError("limits must be a BuyerAttachmentParseLimits instance")
    payload = _as_bytes(content)
    metadata = build_attachment_metadata(payload, filename=filename, content_type=content_type)
    if metadata.size_bytes > active_limits.max_bytes:
        return BuyerAttachmentParseResult(
            metadata=metadata,
            status=BuyerAttachmentParseStatus.TOO_LARGE,
            omissions=(
                BuyerAttachmentOmission(
                    "size_limit",
                    f"{metadata.size_bytes} bytes exceeds the {active_limits.max_bytes}-byte limit",
                ),
            ),
        )
    if metadata.attachment_type is None:
        return BuyerAttachmentParseResult(
            metadata=metadata,
            status=BuyerAttachmentParseStatus.UNSUPPORTED,
            omissions=(BuyerAttachmentOmission("unsupported_type"),),
        )

    attachment_type = metadata.attachment_type
    if attachment_type is BuyerAttachmentType.TEXT:
        text, omissions = _parse_text(payload, active_limits)
        return _parsed_result(metadata, "text", text, omissions, active_limits)
    if attachment_type is BuyerAttachmentType.JSON:
        try:
            text, omissions = _parse_json(payload, active_limits)
        except (TypeError, ValueError, json.JSONDecodeError):
            return _corrupt_result(metadata, "json")
        return _parsed_result(metadata, "json", text, omissions, active_limits)
    if attachment_type is BuyerAttachmentType.CSV:
        try:
            text, omissions = _parse_csv(payload, active_limits)
        except (csv.Error, UnicodeError, ValueError):
            return _corrupt_result(metadata, "csv")
        return _parsed_result(metadata, "csv", text, omissions, active_limits)
    if attachment_type is BuyerAttachmentType.XML:
        try:
            text, omissions = _parse_xml(payload, active_limits)
        except (ImportError, ValueError, UnicodeError):
            return _corrupt_result(metadata, "xml")
        return _parsed_result(metadata, "xml", text, omissions, active_limits)
    return _parse_optional_document(metadata, payload, active_limits)


def build_attachment_context(
    results: Iterable[BuyerAttachmentParseResult],
    *,
    max_total_characters: int = 48_000,
) -> BuyerAttachmentContext:
    """Build a bounded, redacted context and an honest attachment manifest."""

    if isinstance(max_total_characters, bool) or not isinstance(max_total_characters, int) or max_total_characters <= 0:
        raise ValueError("max_total_characters must be a positive integer")

    remaining = max_total_characters
    chunks: list[str] = []
    manifest: list[dict[str, Any]] = []
    for result in results:
        if not isinstance(result, BuyerAttachmentParseResult):
            raise TypeError("results must contain BuyerAttachmentParseResult values")
        entry = result.to_manifest()
        safe_filename = sanitize_attachment_context_text(result.metadata.filename)
        entry["filename"] = safe_filename
        text = sanitize_attachment_context_text(result.text)
        if result.status is not BuyerAttachmentParseStatus.PARSED or not text:
            entry["context_characters"] = 0
            manifest.append(entry)
            continue

        header = f"Attachment: {safe_filename} ({result.metadata.attachment_type.value})\n"
        if remaining <= len(header):
            entry["context_characters"] = 0
            entry["context_omission"] = "context_character_limit"
            manifest.append(entry)
            continue
        available_text = remaining - len(header)
        clipped = text[:available_text]
        chunks.append(f"{header}{clipped}")
        remaining -= len(header) + len(clipped)
        entry["context_characters"] = len(clipped)
        if len(clipped) < len(text):
            entry["context_omission"] = "context_character_limit"
        manifest.append(entry)

    context_text = "\n\n".join(chunks)
    return BuyerAttachmentContext(
        text=context_text,
        manifest=tuple(manifest),
        context_hash=sha256(context_text.encode("utf-8")).hexdigest(),
    )


def sanitize_attachment_context_text(value: str) -> str:
    """Remove common credentials from text before it becomes model context."""

    if not isinstance(value, str):
        raise TypeError("context text must be a string")
    sanitized = _SECRET_HEADER_RE.sub(lambda match: f"{match.group('prefix')} [REDACTED]", value)
    sanitized = _SECRET_ASSIGNMENT_RE.sub(lambda match: f"{match.group('prefix')}[REDACTED]", sanitized)
    sanitized = _BEARER_RE.sub("Bearer [REDACTED]", sanitized)
    return _URL_SECRET_RE.sub(lambda match: f"{match.group('prefix')}[REDACTED]", sanitized)


def _parse_text(payload: bytes, limits: BuyerAttachmentParseLimits) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    decoded = payload.decode("utf-8-sig", errors="replace")
    return _bounded_sanitized_text(decoded, limits)


def _parse_json(payload: bytes, limits: BuyerAttachmentParseLimits) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    decoded = payload.decode("utf-8-sig")
    parsed = json.loads(decoded)
    sanitized = _sanitize_structured_value(parsed)
    rendered = json.dumps(sanitized, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False)
    return _bounded_sanitized_text(rendered, limits)


def _parse_csv(payload: bytes, limits: BuyerAttachmentParseLimits) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    decoded = payload.decode("utf-8-sig", errors="replace")
    rows: list[str] = []
    omissions: list[BuyerAttachmentOmission] = []
    reader = csv.reader(StringIO(decoded, newline=""))
    header: tuple[str, ...] = ()
    for row_index, row in enumerate(reader):
        if row_index >= limits.max_rows:
            omissions.append(BuyerAttachmentOmission("row_limit", f"limited to {limits.max_rows} rows"))
            break
        if len(row) > limits.max_columns:
            omissions.append(BuyerAttachmentOmission("column_limit", f"limited to {limits.max_columns} columns"))
            row = row[: limits.max_columns]
        if row_index == 0:
            header = tuple(row)
            rows.append(" | ".join(sanitize_attachment_context_text(cell) for cell in row))
            continue
        safe_cells = (
            "[REDACTED]" if index < len(header) and _is_sensitive_key(header[index]) else sanitize_attachment_context_text(cell)
            for index, cell in enumerate(row)
        )
        rows.append(" | ".join(safe_cells))
    text, text_omissions = _bounded_sanitized_text("\n".join(rows), limits)
    return text, tuple((*omissions, *text_omissions))


def _parse_optional_document(
    metadata: BuyerAttachmentMetadata,
    payload: bytes,
    limits: BuyerAttachmentParseLimits,
) -> BuyerAttachmentParseResult:
    assert metadata.attachment_type is not None
    parser = {
        BuyerAttachmentType.PDF: _extract_pdf_text,
        BuyerAttachmentType.DOCX: _extract_docx_text,
        BuyerAttachmentType.XLSX: _extract_xlsx_text,
        BuyerAttachmentType.PPTX: _extract_pptx_text,
        BuyerAttachmentType.IMAGE: _extract_image_text,
        BuyerAttachmentType.ZIP: _extract_zip_text,
    }.get(metadata.attachment_type)
    if parser is None:
        return BuyerAttachmentParseResult(
            metadata=metadata,
            status=BuyerAttachmentParseStatus.UNSUPPORTED,
            omissions=(BuyerAttachmentOmission("unsupported_type"),),
        )
    try:
        extracted = parser(payload, limits)
    except ImportError:
        return BuyerAttachmentParseResult(
            metadata=metadata,
            status=BuyerAttachmentParseStatus.UNSUPPORTED,
            parser=metadata.attachment_type.value,
            omissions=(BuyerAttachmentOmission("parser_unavailable"),),
        )
    except _PasswordProtectedDocument:
        return BuyerAttachmentParseResult(
            metadata=metadata,
            status=BuyerAttachmentParseStatus.UNSUPPORTED,
            parser=metadata.attachment_type.value,
            omissions=(BuyerAttachmentOmission("password_protected"),),
        )
    except Exception:  # noqa: BLE001 - malformed third-party document libraries must fail closed.
        return _corrupt_result(metadata, metadata.attachment_type.value)
    text, omissions = extracted
    return _parsed_result(metadata, metadata.attachment_type.value, text, omissions, limits)


def _extract_pdf_text(payload: bytes, limits: BuyerAttachmentParseLimits) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(payload), strict=False)
    if reader.is_encrypted and reader.decrypt("") == 0:
        raise _PasswordProtectedDocument
    pages: list[str] = []
    omissions: list[BuyerAttachmentOmission] = []
    for page_index, page in enumerate(reader.pages):
        if page_index >= limits.max_pages:
            omissions.append(BuyerAttachmentOmission("page_limit", f"limited to {limits.max_pages} pages"))
            break
        pages.append(page.extract_text() or "")
    text, text_omissions = _bounded_sanitized_text("\n".join(pages), limits)
    if not text:
        omissions.append(BuyerAttachmentOmission("no_extractable_text"))
    return text, tuple((*omissions, *text_omissions))


def _extract_docx_text(payload: bytes, limits: BuyerAttachmentParseLimits) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    from docx import Document

    document = Document(BytesIO(payload))
    paragraphs = [paragraph.text for paragraph in document.paragraphs[: limits.max_rows]]
    omissions: list[BuyerAttachmentOmission] = []
    if len(document.paragraphs) > limits.max_rows:
        omissions.append(BuyerAttachmentOmission("row_limit", f"limited to {limits.max_rows} paragraphs"))
    text, text_omissions = _bounded_sanitized_text("\n".join(paragraphs), limits)
    return text, tuple((*omissions, *text_omissions))


def _extract_xlsx_text(payload: bytes, limits: BuyerAttachmentParseLimits) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True, keep_vba=False)
    rows: list[str] = []
    omissions: list[BuyerAttachmentOmission] = []
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets):
            if sheet_index >= limits.max_sheets:
                omissions.append(BuyerAttachmentOmission("sheet_limit", f"limited to {limits.max_sheets} sheets"))
                break
            rows.append(f"Sheet: {worksheet.title}")
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if row_index >= limits.max_rows:
                    omissions.append(BuyerAttachmentOmission("row_limit", f"limited to {limits.max_rows} rows per sheet"))
                    break
                values = tuple("" if value is None else str(value) for value in row[: limits.max_columns])
                if len(row) > limits.max_columns:
                    omissions.append(BuyerAttachmentOmission("column_limit", f"limited to {limits.max_columns} columns"))
                rows.append(" | ".join(sanitize_attachment_context_text(value) for value in values))
    finally:
        workbook.close()
    text, text_omissions = _bounded_sanitized_text("\n".join(rows), limits)
    return text, tuple((*omissions, *text_omissions))


def _extract_pptx_text(payload: bytes, limits: BuyerAttachmentParseLimits) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    from pptx import Presentation

    presentation = Presentation(BytesIO(payload))
    slides: list[str] = []
    omissions: list[BuyerAttachmentOmission] = []
    for slide_index, slide in enumerate(presentation.slides):
        if slide_index >= limits.max_slides:
            omissions.append(BuyerAttachmentOmission("slide_limit", f"limited to {limits.max_slides} slides"))
            break
        text_parts = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text]
        slides.append("\n".join(text_parts))
    text, text_omissions = _bounded_sanitized_text("\n".join(slides), limits)
    return text, tuple((*omissions, *text_omissions))


def _parse_xml(payload: bytes, limits: BuyerAttachmentParseLimits) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    """Extract XML text without permitting document type declarations or entities."""

    if b"<!DOCTYPE" in payload.upper() or b"<!ENTITY" in payload.upper():
        raise ValueError("XML declarations with external entities are not supported")
    from defusedxml import ElementTree

    root = ElementTree.fromstring(payload)
    values: list[str] = []
    omissions: list[BuyerAttachmentOmission] = []
    for index, element in enumerate(root.iter()):
        if index >= limits.max_rows:
            omissions.append(BuyerAttachmentOmission("row_limit", f"limited to {limits.max_rows} XML elements"))
            break
        text = (element.text or "").strip()
        if text:
            tag = str(element.tag).rsplit("}", 1)[-1]
            values.append("[REDACTED]" if _is_sensitive_key(tag) else text)
    bounded, text_omissions = _bounded_sanitized_text("\n".join(values), limits)
    return bounded, tuple((*omissions, *text_omissions))


def _extract_image_text(payload: bytes, limits: BuyerAttachmentParseLimits) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    """Read image metadata only; OCR/vision remains an explicit optional pass."""

    from PIL import Image

    with Image.open(BytesIO(payload)) as image:
        width, height = image.size
        frame_count = int(getattr(image, "n_frames", 1) or 1)
        metadata = f"Image metadata: {width}x{height}; mode={image.mode}; frames={frame_count}"
    bounded, text_omissions = _bounded_sanitized_text(metadata, limits)
    return bounded, (BuyerAttachmentOmission("vision_not_configured"), *text_omissions)


def _extract_zip_text(payload: bytes, limits: BuyerAttachmentParseLimits) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    """Parse supported archive members under strict member and byte ceilings."""

    rows: list[str] = []
    omissions: list[BuyerAttachmentOmission] = []
    total_uncompressed = 0
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        infos = [item for item in archive.infolist() if not item.is_dir()]
        for index, info in enumerate(infos):
            if index >= limits.max_archive_members:
                omissions.append(BuyerAttachmentOmission("archive_member_limit", f"limited to {limits.max_archive_members} members"))
                break
            filename = normalize_attachment_filename(info.filename, fallback="archive-member")
            if info.flag_bits & 0x1:
                omissions.append(BuyerAttachmentOmission("password_protected_member", filename))
                continue
            if info.file_size < 0 or info.file_size > limits.max_bytes or total_uncompressed + info.file_size > limits.max_bytes:
                omissions.append(BuyerAttachmentOmission("archive_uncompressed_size_limit", filename))
                continue
            member_type = detect_attachment_type(filename, None)
            if member_type in {None, BuyerAttachmentType.ZIP}:
                omissions.append(BuyerAttachmentOmission("unsupported_archive_member", filename))
                continue
            total_uncompressed += info.file_size
            try:
                content = archive.read(info)
            except (RuntimeError, zipfile.BadZipFile):
                omissions.append(BuyerAttachmentOmission("corrupt_archive_member", filename))
                continue
            parsed = parse_attachment(content, filename=filename, limits=limits)
            if parsed.status is BuyerAttachmentParseStatus.PARSED and parsed.text:
                rows.append(f"Archive member: {filename}\n{parsed.text}")
            else:
                omissions.append(BuyerAttachmentOmission(f"archive_member_{parsed.status.value}", filename))
            omissions.extend(parsed.omissions)
    bounded, text_omissions = _bounded_sanitized_text("\n\n".join(rows), limits)
    return bounded, tuple((*omissions, *text_omissions))


def _parsed_result(
    metadata: BuyerAttachmentMetadata,
    parser: str,
    text: str,
    omissions: tuple[BuyerAttachmentOmission, ...],
    limits: BuyerAttachmentParseLimits,
) -> BuyerAttachmentParseResult:
    bounded_text, bounded_omissions = _bounded_sanitized_text(text, limits)
    return BuyerAttachmentParseResult(
        metadata=metadata,
        status=BuyerAttachmentParseStatus.PARSED,
        parser=parser,
        text=bounded_text,
        omissions=tuple((*omissions, *bounded_omissions)),
    )


def _corrupt_result(metadata: BuyerAttachmentMetadata, parser: str) -> BuyerAttachmentParseResult:
    return BuyerAttachmentParseResult(
        metadata=metadata,
        status=BuyerAttachmentParseStatus.CORRUPT,
        parser=parser,
        omissions=(BuyerAttachmentOmission("parse_failed"),),
    )


def _bounded_sanitized_text(
    value: str,
    limits: BuyerAttachmentParseLimits,
) -> tuple[str, tuple[BuyerAttachmentOmission, ...]]:
    text = sanitize_attachment_context_text(value)
    if len(text) <= limits.max_characters:
        return text, ()
    return (
        text[: limits.max_characters],
        (BuyerAttachmentOmission("character_limit", f"limited to {limits.max_characters} characters"),),
    )


def _sanitize_structured_value(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {
            str(key): "[REDACTED]" if _is_sensitive_key(str(key)) else _sanitize_structured_value(item)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [_sanitize_structured_value(item) for item in value]
    if isinstance(value, tuple):
        return [_sanitize_structured_value(item) for item in value]
    if isinstance(value, str):
        return sanitize_attachment_context_text(value)
    if isinstance(value, float) and not math.isfinite(value):
        return "[UNSUPPORTED_NUMBER]"
    return value


def _is_sensitive_key(value: str) -> bool:
    normalized = value.strip().casefold().replace(" ", "_")
    if _SENSITIVE_KEY_RE.search(normalized):
        return True
    compacted = normalized.replace("_", "").replace("-", "")
    return any(marker in compacted for marker in ("authorization", "cookie", "password", "secret", "token", "apikey", "session"))


def _filename_parts(value: str) -> tuple[str, str]:
    index = value.rfind(".")
    if index <= 0 or index == len(value) - 1:
        return value, ""
    return value[:index], value[index:]


def _as_bytes(value: bytes | bytearray | memoryview) -> bytes:
    if not isinstance(value, (bytes, bytearray, memoryview)):
        raise TypeError("attachment content must be bytes-like")
    return bytes(value)


class _PasswordProtectedDocument(Exception):
    """Internal signal for a document parser that must not attempt a password."""


__all__ = [
    "BuyerAttachmentContext",
    "BuyerAttachmentMetadata",
    "BuyerAttachmentOmission",
    "BuyerAttachmentParseLimits",
    "BuyerAttachmentParseResult",
    "BuyerAttachmentParseStatus",
    "BuyerAttachmentType",
    "build_attachment_context",
    "build_attachment_metadata",
    "detect_attachment_type",
    "normalize_attachment_content_type",
    "normalize_attachment_filename",
    "parse_attachment",
    "sanitize_attachment_context_text",
]
